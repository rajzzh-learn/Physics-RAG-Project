"""
Utility helpers for processing file attachments uploaded via the Streamlit UI.

Supported attachment types:
  - Images (PNG / JPG / JPEG / WEBP / GIF) -> base64 data URI for vision LLMs
  - PDF            -> text via pypdf
  - DOCX           -> paragraphs + tables via python-docx
  - PPTX           -> slide text + tables via python-pptx
  - XLSX           -> cell values via openpyxl
  - CSV / TXT / MD / JSON / TEX -> decoded text
  - HTML / HTM     -> visible text (scripts/styles stripped)
"""
from __future__ import annotations

import base64
import csv
import io
import re
from html.parser import HTMLParser
from typing import Tuple

DEFAULT_MAX_CHARS = 6000

# -- Helpers ----------------------------------------------------------------

def _truncate(text: str, max_chars: int) -> str:
    text = text.strip()
    if len(text) > max_chars:
        text = text[:max_chars] + f"\n\n[truncated at {max_chars} chars]"
    return text


def _decode(file_bytes: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return file_bytes.decode(enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return file_bytes.decode("utf-8", errors="replace")


# -- PDF --------------------------------------------------------------------

def extract_text_from_pdf(file_bytes: bytes, max_chars: int = DEFAULT_MAX_CHARS) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(file_bytes))
    pages_text = [page.extract_text() or "" for page in reader.pages]
    return _truncate("\n\n".join(t.strip() for t in pages_text if t.strip()), max_chars)


# -- Plain text -------------------------------------------------------------

def extract_text_from_txt(file_bytes: bytes, max_chars: int = DEFAULT_MAX_CHARS) -> str:
    return _truncate(_decode(file_bytes), max_chars)


# -- DOCX -------------------------------------------------------------------

def extract_text_from_docx(file_bytes: bytes, max_chars: int = DEFAULT_MAX_CHARS) -> str:
    from docx import Document
    doc = Document(io.BytesIO(file_bytes))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return _truncate("\n".join(lines), max_chars)


# -- PPTX -------------------------------------------------------------------

def extract_text_from_pptx(file_bytes: bytes, max_chars: int = DEFAULT_MAX_CHARS) -> str:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(file_bytes))
    out: list[str] = []
    for idx, slide in enumerate(prs.slides, start=1):
        out.append(f"--- Slide {idx} ---")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    out.append(" | ".join(c.text.strip() for c in row.cells))
    return _truncate("\n".join(out), max_chars)


# -- XLSX / CSV -------------------------------------------------------------

def extract_text_from_xlsx(file_bytes: bytes, max_chars: int = DEFAULT_MAX_CHARS) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    out: list[str] = []
    for ws in wb.worksheets:
        out.append(f"--- Sheet: {ws.title} ---")
        for row in ws.iter_rows(values_only=True):
            if any(v is not None for v in row):
                out.append(" | ".join("" if v is None else str(v) for v in row))
            if sum(len(s) for s in out) > max_chars:
                break
    wb.close()
    return _truncate("\n".join(out), max_chars)


def extract_text_from_csv(file_bytes: bytes, max_chars: int = DEFAULT_MAX_CHARS) -> str:
    rows = csv.reader(io.StringIO(_decode(file_bytes)))
    return _truncate("\n".join(" | ".join(r) for r in rows), max_chars)


# -- HTML -------------------------------------------------------------------

class _TextExtractor(HTMLParser):
    _SKIP = {"script", "style", "head", "noscript", "template"}
    _BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "table", "section"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag, attrs):
        if tag in self._SKIP:
            self._skip_depth += 1
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self._SKIP and self._skip_depth:
            self._skip_depth -= 1
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self._skip_depth:
            self.parts.append(data)


def extract_text_from_html(file_bytes: bytes, max_chars: int = DEFAULT_MAX_CHARS) -> str:
    parser = _TextExtractor()
    parser.feed(_decode(file_bytes))
    text = re.sub(r"[ \t]+", " ", "".join(parser.parts))
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return _truncate(text, max_chars)


# -- Dispatch ---------------------------------------------------------------

_EXTRACTORS = {
    "pdf": extract_text_from_pdf,
    "docx": extract_text_from_docx,
    "pptx": extract_text_from_pptx,
    "xlsx": extract_text_from_xlsx,
    "csv": extract_text_from_csv,
    "html": extract_text_from_html,
    "htm": extract_text_from_html,
    "txt": extract_text_from_txt,
    "md": extract_text_from_txt,
    "markdown": extract_text_from_txt,
    "json": extract_text_from_txt,
    "tex": extract_text_from_txt,
}


def extract_text(file_bytes: bytes, filename: str, max_chars: int = DEFAULT_MAX_CHARS) -> str:
    """Extract text from any supported non-image attachment."""
    ext = _ext(filename)
    fn = _EXTRACTORS.get(ext)
    if fn is None:
        raise ValueError(f"Unsupported file type: .{ext}")
    return fn(file_bytes, max_chars=max_chars)


# -- Image -> base64 data URI -----------------------------------------------

_MIME_MAP = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg",
             "webp": "image/webp", "gif": "image/gif"}

SUPPORTED_IMAGE_EXTS = set(_MIME_MAP.keys())
SUPPORTED_DOC_EXTS   = set(_EXTRACTORS.keys())
SUPPORTED_EXTS       = SUPPORTED_IMAGE_EXTS | SUPPORTED_DOC_EXTS


def _ext(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def image_to_base64_uri(file_bytes: bytes, filename: str) -> Tuple[str, str]:
    mime = _MIME_MAP.get(_ext(filename), "image/png")
    b64  = base64.b64encode(file_bytes).decode("ascii")
    return f"data:{mime};base64,{b64}", mime


def is_image(filename: str) -> bool:
    return _ext(filename) in SUPPORTED_IMAGE_EXTS
